"""通用 Markdown → Word(docx)转换器(项目文档统一排版)。

用法: python scripts/md_to_docx.py <输入.md> [输出.docx]
支持:标题、段落、无序/有序列表、引用块、表格(含表头加粗)、水平线忽略、行内 **加粗**/`代码` 标记剥离。
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt


def _plain(text: str) -> str:
    return text.replace("**", "").replace("`", "").strip()


def convert(md_path: Path, out_path: Path) -> int:
    lines = md_path.read_text(encoding="utf-8").splitlines()
    doc = docx.Document()
    style = doc.styles["Normal"]
    style.font.name = "微软雅黑"
    style.font.size = Pt(10.5)
    try:
        style.element.rPr.rFonts.set(docx.oxml.ns.qn("w:eastAsia"), "微软雅黑")
    except Exception:
        pass

    i, tables = 0, 0
    while i < len(lines):
        s = lines[i].strip()
        if s.startswith("|"):
            block = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                block.append(lines[i].strip())
                i += 1
            rows = []
            for r in block:
                if set(r.replace("|", "").strip()) <= {"-", " "}:
                    continue
                rows.append([c.strip() for c in r.strip("|").split("|")])
            if rows:
                t = doc.add_table(rows=len(rows), cols=max(len(r) for r in rows))
                t.style = "Light Grid Accent 1"
                for ri, r in enumerate(rows):
                    for ci, cell in enumerate(r):
                        if ci < len(t.rows[ri].cells):
                            run = t.rows[ri].cells[ci].paragraphs[0].add_run(_plain(cell))
                            run.font.size = Pt(9)
                            if ri == 0:
                                run.bold = True
                doc.add_paragraph()
                tables += 1
            continue
        if s.startswith("# "):
            h = doc.add_heading(_plain(s[2:]), level=0)
            h.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif s.startswith("## "):
            doc.add_heading(_plain(s[3:]), level=1)
        elif s.startswith("### "):
            doc.add_heading(_plain(s[4:]), level=2)
        elif s.startswith("#### "):
            doc.add_heading(_plain(s[5:]), level=3)
        elif s.startswith("> "):
            p = doc.add_paragraph(_plain(s[2:]))
            if p.runs:
                p.runs[0].italic = True
        elif s.startswith("- ") or s.startswith("* "):
            doc.add_paragraph(_plain(s[2:]), style="List Bullet")
        elif re.match(r"^\d+\.\s", s):
            doc.add_paragraph(_plain(re.sub(r"^\d+\.\s", "", s)), style="List Number")
        elif s and s != "---":
            doc.add_paragraph(_plain(s))
        i += 1

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    return tables


def main() -> int:
    if len(sys.argv) < 2:
        print("用法: python scripts/md_to_docx.py <输入.md> [输出.docx]")
        return 2
    src = Path(sys.argv[1])
    if not src.exists():
        print(f"文件不存在: {src}")
        return 1
    dst = Path(sys.argv[2]) if len(sys.argv) > 2 else src.with_suffix(".docx")
    tables = convert(src, dst)
    print(f"已生成:{dst}(段落 {len(docx.Document(str(dst)).paragraphs)} 个,表格 {tables} 个)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
