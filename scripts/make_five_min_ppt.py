"""生成《答辩 PPT(5 分钟精简版)》:14 页,与 5 分钟讲稿页号一一对应。

设计原则:5 分钟 ≈ 300 秒,平均每页 20 秒;文字量压缩、字号放大,一页只讲一件事。
用法: python scripts/make_five_min_ppt.py
"""
from __future__ import annotations

import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import make_defense_ppt as kit  # 复用主题与版式工具
from pptx.util import Inches

DOCS = kit.DOCS
OUT = DOCS / "答辩PPT(5分钟精简版).pptx"
W = kit.W


def build():
    prs = kit.Presentation()
    prs.slide_width = kit.W
    prs.slide_height = kit.H
    blank, header, footer = kit.blank, kit.header, kit.footer
    bullets, cards, table, highlight = kit.bullets, kit.cards, kit.table, kit.highlight
    rect, textbox, para = kit.rect, kit.textbox, kit.para
    NAVY, BLUE, TEAL, AMBER, GREY, WHITE = kit.NAVY, kit.BLUE, kit.TEAL, kit.AMBER, kit.GREY, kit.WHITE

    # ===== 1 封面(0:00–0:15) =====
    s = blank(prs)
    rect(s, Inches(0), Inches(0), kit.W, kit.H, NAVY)
    rect(s, Inches(0), Inches(2.52), kit.W, Inches(0.06), TEAL)
    _t, tf = textbox(s, Inches(1.0), Inches(1.05), Inches(11.3), Inches(1.5))
    para(tf, "面向中小企业内部制度文档的", size=22, color=kit.RGBColor(0x9F, 0xC6, 0xE6),
         first=True, space_after=4)
    para(tf, "多知识库私有 RAG 智能问答系统", size=40, bold=True, color=WHITE, space_after=0)
    _t2, tf2 = textbox(s, Inches(1.0), Inches(2.9), Inches(11.3), Inches(1.5))
    para(tf2, "本地私有化 · 库级物理隔离 · 制度版本治理 · 冲突智能识别", size=17,
         color=kit.RGBColor(0xC9, 0xD8, 0xE8), first=True, space_after=10)
    para(tf2, "AIC 全国大学生人工智能创新大赛 · 项目答辩(5 分钟汇报)", size=15, bold=True,
         color=kit.RGBColor(0x7F, 0xD1, 0xC1), space_after=0)
    _t3, tf3 = textbox(s, Inches(1.0), Inches(5.4), Inches(11.3), Inches(0.9))
    para(tf3, "答辩团队:智启未来 AI 团队", size=14, color=WHITE, first=True)

    # ===== 2 痛点(0:15–0:55) =====
    s = blank(prs)
    header(s, "01", "问题:制度文档管不住,通用方案不敢用", "中小企业的制度痛点 × 传统 RAG 的四个失效点")
    cards(s, [
        ("困境一:信息孤岛", ["制度散落在电脑/邮件/群聊/共享盘", "找文件靠问同事,新人上手慢"], AMBER),
        ("困境二:版本混乱", ["试行版/修订版/废止版并存", "员工不知道哪一版算数,执行偏差"], AMBER),
        ("困境三:查询低效", ["关键词搜索大量无关结果", "条款级信息(天数/期限)容易漏查"], AMBER),
        ("痛点一:隔离不足", ["所有文档进同一向量库", "跨部门跨版本相互污染"], kit.RGBColor(0xB0, 0x3A, 0x2B)),
        ("痛点二:更新昂贵", ["改一处要整库重建索引", "制度迭代快,索引长期滞后"], kit.RGBColor(0xB0, 0x3A, 0x2B)),
        ("痛点三/四:冲突与泄露", ["新旧矛盾被模型“融合”成错误答案", "内部制度上传公有云,合规风险"], kit.RGBColor(0xB0, 0x3A, 0x2B)),
    ], top=Inches(1.5), height=Inches(4.6), per_row=3)
    highlight(s, "一句话:我们要做的不是“更聪明的问答”,而是“口径唯一、数据不出内网”的制度治理工具", top=Inches(6.25))
    footer(s)

    # ===== 3 建设目标(0:55–1:25) =====
    s = blank(prs)
    header(s, "02", "建设目标:四条,逐条可验证", "每项目标都对应系统机制与验证断言")
    cards(s, [
        ("① 物理隔离", ["每个知识库独立 Chroma 实例", "独立目录、独立 sqlite 文件"], TEAL),
        ("② 增量维护", ["块哈希账本,只重算变更块", "删除按文档 ID 精准清理"], BLUE),
        ("③ 冲突识别", ["多库检索 + 双通道判定", "版本差异单独处置,以现行版为准"], NAVY),
        ("④ 全本地化", ["Ollama 本地推理,零外部 API", "数据不出内网,断网可用"], kit.RGBColor(0x6A, 0x4C, 0x9B)),
    ], top=Inches(1.55), height=Inches(4.2), per_row=4)
    highlight(s, "额外交付:部门权限隔离 + 审计留痕,让制度系统“可控可见、可追溯”", top=Inches(6.0), color=BLUE)
    footer(s)

    # ===== 4 架构(1:25–1:50) =====
    s = blank(prs)
    header(s, "02", "总体架构:轻量三层 + 本地模型层", "模块解耦,普通办公电脑即可部署")
    rows = [
        ("应用层", "FastAPI 托管的原生 Web 界面:知识库 / 文档 / 问答 / 自动分类建库 / 审计 / 系统状态", BLUE),
        ("服务层", "制度与文档治理(元数据·版本·增量索引)· 多库检索路由 · 冲突与版本差异检测 · 权限 · 审计", TEAL),
        ("数据层", "每库独立 Chroma 实例(data/vectors/<kb_id>/)+ SQLite(库/文档/制度元数据/块账本/审计)", NAVY),
        ("模型层", "Ollama 本地:deepseek-r1:7b(生成,4.7GB)+ bge-m3(嵌入,1024 维)", kit.RGBColor(0x6A, 0x4C, 0x9B)),
    ]
    top = Inches(1.5)
    for name, desc, color in rows:
        box = rect(s, Inches(0.7), top, Inches(11.93), Inches(1.0), kit.LIGHT)
        tf = box.text_frame
        tf.margin_left = Inches(0.24)
        para(tf, name, size=14.5, bold=True, color=color, first=True, space_after=2)
        para(tf, desc, size=11.5, color=GREY, space_after=0)
        top += Inches(1.1)
    highlight(s, "亮点:库级物理隔离 · 版本可治理 · 冲突可识别 · 权限可管控 · 全流程可审计 · 零外部依赖", top=Inches(6.05))
    footer(s)

    # ===== 5 能力一(1:50–2:15) =====
    s = blank(prs)
    header(s, "03", "核心能力一:多知识库向量硬隔离", "从“打标签的逻辑隔离”到“实例级物理隔离”")
    bullets(s, [
        "每个知识库 = 独立 ChromaDB 实例 + 独立存储目录 + 独立 chroma.sqlite3;",
        "删库即物理删除该库目录,不触碰其它库;旧布局启动时自动迁移并校验数量(保留备份可回滚);",
        "隔离证据可查:输出每库存储路径、独立 sqlite、文件数与磁盘占用。",
    ], top=Inches(1.5), size=15)
    table(s, ["实测", "结果"],
          [["删除 A 库", "A 目录物理移除(mode=deleted),B 库文件完好、问答检索正常"],
           ["独立路径数", "独立路径数 = 知识库数(每库 1 实例 1 sqlite)"]],
          top=Inches(3.35), height=Inches(1.5), fsize=12)
    highlight(s, "解决:通用方案“单库混杂”导致的信息污染与来源无法界定", top=Inches(5.4))
    footer(s)

    # ===== 6 能力二(2:15–2:40) =====
    s = blank(prs)
    header(s, "03", "核心能力二:块级增量更新 + 精准删除", "从“整篇重向量化”到“按改动量更新”")
    bullets(s, [
        "块哈希账本:切块后逐块计算内容指纹,建立“块–哈希–向量”映射;",
        "未变更块直接复用既有向量,仅对改动块重新嵌入;删除按文档 ID 精准清理并二次校验。",
    ], top=Inches(1.5), size=15)
    table(s, ["指标", "实测值"],
          [["复用未变更块", "1 / 2 块(复用率 50%)— 5 段文档仅改 1 段"],
           ["更新耗时", "637 ms(整篇重嵌入估算约 1250 ms)"],
           ["删除后残留", "0(zero_residue = true,二次校验 + 对账)"]],
          top=Inches(2.95), height=Inches(2.1), fsize=12)
    highlight(s, "价值:更新成本随“改动量”增长,而非随“库规模”增长", top=Inches(5.55), color=BLUE)
    footer(s)

    # ===== 7 能力三(2:40–3:05) =====
    s = blank(prs)
    header(s, "03", "核心能力三:跨知识库信息冲突检测", "不让模型把矛盾规定“融合”成通顺的错误答案")
    cards(s, [
        ("① 多库检索", ["同一问题在多库并行召回", "问题向量只嵌入一次"], BLUE),
        ("② 话题初筛", ["相似度筛出“同一件事”的片段对", "数值分歧对优先排序"], TEAL),
        ("③ 双通道判定", ["模型通道:语义矛盾判断", "数值通道:取值互斥确定性判定"], NAVY),
        ("④ 告警与约束", ["界面弹出冲突面板", "回答分列来源,禁止擅自合并"], kit.RGBColor(0x6A, 0x4C, 0x9B)),
    ], top=Inches(1.5), height=Inches(3.5), per_row=4)
    highlight(s, "实测:新旧年假制度场景检出冲突,判定通道 numeric+llm;回答并列呈现两库说法并标注来源",
              top=Inches(5.3), size=12.5)
    footer(s)

    # ===== 8 能力四(3:05–3:35) =====
    s = blank(prs)
    header(s, "03", "核心能力四:制度版本与生效期治理", "把“冲突”升级为可判定的“版本问题”")
    cards(s, [
        ("制度元数据", ["文件编号 / 生效日期 / 失效日期", "状态:现行 · 已被替代 · 草案 · 已失效", "发布部门 / 适用范围 / 替代关系"], TEAL),
        ("检索侧治理", ["自动排除已废止与未生效版本", "同制度多版本时以现行版为依据", "实测:命中仅 2025 版,回答按现行口径"], BLUE),
        ("版本差异专类", ["同制度不同版本 → 判为「版本差异」", "给出建议依据的现行版本编号", "不同来源矛盾 → 判为「跨来源冲突」"], NAVY),
        ("版本链与提醒", ["版本链:按生效日期列出全部版本", "到期提醒:30 天内到期清单", "支撑制度滚动复审"], kit.RGBColor(0x6A, 0x4C, 0x9B)),
    ], top=Inches(1.5), height=Inches(3.9), per_row=2)
    highlight(s, "实测:新版声明替代后旧版自动“已被替代”且不再被检索;跨库版本差异给出 preferred=HR-2025-777",
              top=Inches(5.6), size=12.5, color=BLUE)
    footer(s)

    # ===== 9 能力五六(3:35–3:55) =====
    s = blank(prs)
    header(s, "03", "核心能力五 / 六:权限隔离与审计留痕", "敏感制度可控可见,操作全程可追溯")
    cards(s, [
        ("五 · 部门 / 角色权限", ["知识库可见范围:全员 或 指定部门",
                               "库列表 / 路由 / 检索 / 冲突检测 / 文档操作全链路鉴权",
                               "越权:HTTP 403 + 明确拒绝原因 + 审计留痕",
                               "实测:财务部访问受限薪酬库被拒;HR 正常问答"], BLUE),
        ("六 · 审计日志", ["九类操作留痕:问答 / 冲突 / 越权 / 文档增删改 / 建库等",
                         "可按操作类型、知识库、部门、结果过滤",
                         "统计:总数、越权次数、操作分布、最近时间",
                         "非管理员仅能查询本部门记录"], TEAL),
    ], top=Inches(1.6), height=Inches(4.0), per_row=2)
    highlight(s, "价值:一台内网设备即实现“部门级制度隔离 + 合规留痕”,无需额外权限系统", top=Inches(5.9))
    footer(s)

    # ===== 10 稳定性(3:55–4:15) =====
    s = blank(prs)
    header(s, "04", "稳定性工程:让系统长期跑得住", "本地大模型的真实瓶颈:显存、并发、进程可靠性")
    cards(s, [
        ("进程可靠性", ["后端崩溃 5 秒自愈(守护脚本)", "Ollama 宕机自动唤醒", "实测:强杀后端→5s 恢复;杀 Ollama→6s 唤醒"], TEAL),
        ("显存与并发", ["推理全局串行:1 嵌入 + 1 对话", "8.9GB 模型 → 4.7GB,消除 OOM 主因", "限制同时加载模型数与驻留时长"], BLUE),
        ("请求韧性", ["重 IO 端点线程池化,不阻塞全站", "调用超时 + 前置探活,失败快速提示", "前端 SSE 总超时 + 空闲中断 + 停止按钮"], NAVY),
        ("检索可靠性", ["低相关片段拒答,不硬编", "条款级切分避免跨主题串味", "对账清理孤儿向量、校验账本"], kit.RGBColor(0x6A, 0x4C, 0x9B)),
    ], top=Inches(1.55), height=Inches(4.1), per_row=2)
    highlight(s, "工程原则:能用确定性方法保证的,不交给模型“猜”", top=Inches(5.9), color=BLUE)
    footer(s)

    # ===== 11 验证体系(4:15–4:30) =====
    s = blank(prs)
    header(s, "05", "实验验证:73 项断言全部通过", "三套脚本化测试,一条命令可复现,证据 JSON 存档")
    table(s, ["验证套件", "断言", "覆盖内容", "结果"],
          [["产品化增强验证", "27", "制度版本治理 11 · 权限隔离 8 · 审计 8", "全部通过"],
           ["三项创新验证", "27", "硬隔离 7 · 增量更新与删除 12 · 冲突检测 8", "全部通过"],
           ["端到端功能验收", "19", "建库 / 入库 / 更新 / 删除 / 路由 / 问答 / 自动分类", "全部通过"]],
          top=Inches(1.6), height=Inches(2.4), fsize=13)
    highlight(s, "合计 73 项断言 100% 通过;另完成两次故障恢复演练(后端自愈、模型唤醒)",
              top=Inches(4.5), size=14)
    footer(s)

    # ===== 12 实测数据(4:30–4:50) =====
    s = blank(prs)
    header(s, "05", "关键实测数据", "回答“到底好在哪里”,每个数字都有测量口径")
    table(s, ["指标", "实测值", "口径"],
          [["增量更新复用率", "50%(2 块复用 1 块)", "长文档仅改 1 段"],
           ["更新耗时", "637 ms / 全量约 1250 ms", "实际 vs 全量重嵌入估算"],
           ["删除残留向量", "0", "按文档 ID 删除后二次校验 + 对账"],
           ["冲突检测", "检出冲突并分类正确", "新旧年假制度场景,通道 numeric+llm"],
           ["权限拦截", "HTTP 403 + 拒绝事件", "非授权部门访问受限库"],
           ["单次问答响应", "约 13–19 秒", "deepseek-r1:7b,8GB 显存笔记本"]],
          top=Inches(1.55), height=Inches(3.9), fsize=12)
    footer(s)

    # ===== 13 场景闭环(弹性页 4:50–5:00) =====
    s = blank(prs)
    header(s, "05", "场景闭环:制度换版,口径自动统一", "以“新旧年假制度”为例(可现场演示)")
    bullets(s, [
        "① 提问“年假几天?未休完能否结转?” → ② 多库检索并初筛出同一话题片段对;",
        "③ 数值通道发现 5/3 天 与 8/6 月互斥,模型通道确认语义矛盾;",
        "④ 判定为「制度版本差异」,建议依据 HR-2025-777(现行有效);",
        "⑤ 回答分列两版口径与来源,提示以现行版为准;若新版已声明替代,旧版直接不再被检索;",
        "⑥ 全流程写入审计日志,可事后追溯。",
    ], top=Inches(1.5), size=14.5)
    highlight(s, "价值:制度换版不再依赖“通知到人”,系统在回答层面自动执行“口径唯一”", top=Inches(5.5), color=BLUE)
    footer(s)

    # ===== 14 总结(5:00 收尾) =====
    s = blank(prs)
    header(s, "06", "总结与展望", "一个可运行、可验证、可长期维护的制度问答系统")
    cards(s, [
        ("已交付", ["四项目标全部落地并验证",
                   "三项技术创新 + 三项产品化能力",
                   "73 项断言全通过,故障恢复演练有效",
                   "8GB 显存普通笔记本稳定运行"], TEAL),
        ("下一步", ["制度治理深化:滚动复审提醒、条款级定位",
                   "组织架构对接:账号映射、条款级权限",
                   "性能与规模:重排优化、量化与并发",
                   "交付标准化:一键部署包、备份演练、监控"], BLUE),
    ], top=Inches(1.6), height=Inches(4.0), per_row=2)
    highlight(s, "用确定性工程兜住大模型的不确定性 —— 让制度问答在企业内网真正可用、可控、可追溯",
              top=Inches(5.9), color=NAVY)
    footer(s)
    return prs


def main() -> int:
    DOCS.mkdir(parents=True, exist_ok=True)
    prs = build()
    prs.save(str(OUT))
    print(f"已生成:{OUT}")
    print(f"幻灯片总数:{len(prs.slides._sldIdLst)}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
